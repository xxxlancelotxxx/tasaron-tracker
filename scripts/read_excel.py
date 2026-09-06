import openpyxl
import sys

path = r"C:\Users\berkant.senturk\Desktop\teknikofis\GES_Nagatino-2_Hotel_Building_Follow-up_Report_20260705.xlsm"
print("Dosya aciliyor (makro-ozellikli, buyuk):", path)

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("\n=== SHEET (SAYFA) LISTESI ===")
for i, name in enumerate(wb.sheetnames):
    ws = wb[name]
    print(f"{i+1}. {name}  (satir: {ws.max_row}, kolon: {ws.max_column})")
