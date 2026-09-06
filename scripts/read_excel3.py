import openpyxl
import io

path = r"C:\Users\berkant.senturk\Desktop\teknikofis\GES_Nagatino-2_Hotel_Building_Follow-up_Report_20260705.xlsm"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

out = io.StringIO()

def dump_sheet(out, name, max_rows=22, max_cols=40):
    if name not in wb.sheetnames:
        out.write(f"\n[YOK] {name}\n")
        return
    ws = wb[name]
    out.write(f"\n{'='*90}\nSHEET: {name}  ({ws.max_row} x {ws.max_column})\n{'='*90}\n")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True)):
        vals = [str(c) if c is not None else '' for c in row]
        line = ' | '.join(vals).rstrip(' |')
        if line.strip():
            out.write(f"R{i+1}: {line}\n")

for s in ["Spent MH", "Budget Code Detailed", "Weekly Personnel", "3.Progress Summary", "4.Evaluation", "6. Claims and Change Orders"]:
    dump_sheet(out, s)

with open(r"C:\Users\berkant.senturk\AppData\Local\Temp\opencode\excel_dump.txt", "w", encoding="utf-8") as f:
    f.write(out.getvalue())
print("Yazildi. Boyut:", len(out.getvalue()), "karakter")
